"""数据导出测试：/export/{kind} CSV/xlsx。

策略同 M3：临时 SQLite + TestClient。
运行：python -m pytest tests/test_m9_export.py -v
"""
import importlib
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import app.config as cfg_module

VALID_KEY = "test-key-xyz"


class TestExport(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        cfg_module.DB_PATH = os.path.join(self.tmp, "test_m9.db")
        cfg_module.API_KEY = VALID_KEY
        import app.db as db_module
        importlib.reload(db_module)
        db_module.init_db()
        import app.service.tasks as tasks_module
        importlib.reload(tasks_module)
        # 先建任务（products.task_id 有外键约束），再塞两条 products
        self.tid = tasks_module.create_task("detail", {"ids": ["111", "222"]})
        with db_module.get_conn() as conn:
            conn.execute(
                "INSERT INTO products (task_id, product_id, title, price) VALUES (?,?,?,?)",
                (self.tid, "111", "中文标题 A", 9.9),
            )
            conn.execute(
                "INSERT INTO products (task_id, product_id, title, price) VALUES (?,?,?,?)",
                (self.tid, "222", "Title B", 19.9),
            )
        import app.api as api_module
        importlib.reload(api_module)
        from fastapi.testclient import TestClient
        self.client = TestClient(api_module.app)

    def _h(self):
        return {"X-API-Key": VALID_KEY}

    def test_csv_export(self):
        r = self.client.get("/export/products?fmt=csv&task_id=" + str(self.tid) + "", headers=self._h())
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/csv", r.headers["content-type"])
        self.assertIn("attachment", r.headers["content-disposition"])
        text = r.content.decode("utf-8-sig")
        self.assertIn("product_id", text)       # 表头
        self.assertIn("111", text)
        self.assertIn("中文标题 A", text)        # 中文正常

    def test_xlsx_export(self):
        r = self.client.get("/export/products?fmt=xlsx&task_id=" + str(self.tid) + "", headers=self._h())
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        # 能被 openpyxl 打开且有 3 行（表头 + 2 数据）
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        self.assertEqual(ws.max_row, 3)

    def test_listings_empty_ok(self):
        r = self.client.get("/export/listings?fmt=csv", headers=self._h())
        self.assertEqual(r.status_code, 200)  # 空也返回（仅 BOM/空内容）

    def test_bad_kind_400(self):
        self.assertEqual(self.client.get("/export/foo?fmt=csv", headers=self._h()).status_code, 400)

    def test_bad_fmt_400(self):
        self.assertEqual(self.client.get("/export/products?fmt=pdf", headers=self._h()).status_code, 400)

    def test_no_key_401(self):
        self.assertEqual(self.client.get("/export/products?fmt=csv").status_code, 401)


if __name__ == "__main__":
    unittest.main()
